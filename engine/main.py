# -*- coding: utf-8 -*-
"""引擎入口：stdin/stdout JSON-RPC sidecar（供 Electron 调用）+ 命令行模式。

协议（每行一个 JSON 对象）：
  请求:  {"id": 1, "method": "process", "params": {...}}
  响应:  {"id": 1, "result": ...} 或 {"id": 1, "error": "..."}
  通知:  {"method": "progress", "params": {"stage": "...", "percent": 0.5, "message": "..."}}

方法：
  ping        → {ok, version}
  parse_text  {text} → 材料分组预览
  process     {text, audio_path, out_dir, title?, model_size?, use_whisperx?, language?, min_match?} → 完整流水线
  cancel      → 取消正在运行的 process

CLI：python main.py --cli --text paper.txt --audio listening.mp3 --out out_dir [--model small]
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import threading
import traceback
from typing import Any, Callable, Dict, Optional

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from parser import parse_paper_text
from asr_align import ASREngine
from matcher import align_materials, MaterialSpan, SentenceSpan
from cutter import export_courseware
from ai_parser import ai_parse_materials, ai_test_connection
from pdf_reader import extract_pdf_text, looks_like_garbage

VERSION = "0.1.0"


def read_text_source(path_or_text: str, progress_cb=None) -> str:
    """支持直接文本、.txt / .pdf / .docx 路径（懒加载对应库）。

    PDF 在文字层乱码或扫描时自动 OCR；progress_cb(stage, done, total) 可选。
    """
    if os.path.isfile(path_or_text):
        lower = path_or_text.lower()
        if lower.endswith(".pdf"):
            from pdf_reader import extract_pdf_text

            def _pcb(stage, done, total):
                if progress_cb:
                    progress_cb(stage, done, total)

            text, _ocr_used = extract_pdf_text(path_or_text, progress_cb=_pcb)
            return text
        if lower.endswith(".docx"):
            try:
                import docx  # python-docx，可选依赖
            except ImportError as e:
                raise RuntimeError("解析 .docx 需要安装 python-docx：pip install python-docx") from e
            doc = docx.Document(path_or_text)
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        with open(path_or_text, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    return path_or_text


class Sidecar:
    def __init__(self, send: Callable[[dict], None]):
        self._asr: Optional[ASREngine] = None
        self._asr_model_size: Optional[str] = None
        self._cancel = threading.Event()
        self._send = send

    def _progress(self, stage: str, percent: float, message: str = ""):
        self._send({"method": "progress", "params": {
            "stage": stage, "percent": round(max(0.0, min(1.0, percent)), 4), "message": message,
        }})

    def _ensure_asr(self, model_size: str) -> ASREngine:
        if self._asr is None or self._asr_model_size != model_size:
            self._progress("model", 0.0, f"加载模型 {model_size}（首次需下载）")
            self._asr = ASREngine(
                model_size=model_size,
                model_dir=os.environ.get("YP_MODEL_DIR"),
                device="cpu",
                compute_type="int8",
            )
            self._asr_model_size = model_size
        return self._asr

    # ---- RPC 方法 ----
    def ping(self) -> dict:
        return {"ok": True, "version": VERSION}

    def cancel(self) -> dict:
        self._cancel.set()
        return {"ok": True}

    @staticmethod
    def _dump_materials(materials) -> list:
        return [
            {
                "id": m.id, "title": m.title, "text": m.text,
                "sentences": [{"text": s.text, "speaker": s.speaker} for s in m.sentences],
            }
            for m in materials
        ]

    def parse_text(self, text: str) -> dict:
        materials = parse_paper_text(text)
        return {"materials": self._dump_materials(materials)}

    def ai_test(self, ai_api_key: str, ai_base_url: str, ai_model: str = "") -> dict:
        """测试 AI 接口连接与 Key 有效性。"""
        return ai_test_connection(ai_base_url, ai_api_key, ai_model)

    def read_text(self, path: str) -> dict:
        """读取文稿文件内容用于预览（扫描/乱码 PDF 自动 OCR）。"""
        if not os.path.isfile(path):
            raise RuntimeError(f"文件不存在: {path}")
        lower = path.lower()
        if lower.endswith(".pdf"):
            def _pcb(stage, done, total):
                self._progress("ocr", done / max(1, total), f"OCR 识别第 {done + 1}/{total} 页")
            text, ocr_used = extract_pdf_text(path, progress_cb=_pcb)
            return {"text": text, "is_scanned": False, "garbage": False, "ocr_used": ocr_used}
        elif lower.endswith(".docx"):
            try:
                import docx
            except ImportError as e:
                raise RuntimeError("解析 .docx 需要安装 python-docx") from e
            doc = docx.Document(path)
            text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        return {"text": text, "is_scanned": False, "garbage": looks_like_garbage(text),
                "ocr_used": False}

    @staticmethod
    def _materials_from_dicts(items: list) -> list:
        """把前端校对后的材料 JSON 构造成 parser.Material。"""
        from parser import Material, Sentence
        out: list = []
        for i, m in enumerate(items or [], 1):
            sents = []
            for s in m.get("sentences", []) or []:
                t = (s.get("text") or "").strip()
                if t:
                    sp = s.get("speaker")
                    sp = sp.strip() if isinstance(sp, str) and sp.strip() else None
                    sents.append(Sentence(text=t, speaker=sp))
            if not sents:
                continue
            out.append(Material(
                id=len(out) + 1,
                title=(m.get("title") or f"Text {i}").strip(),
                text=" ".join(s.text for s in sents),
                sentences=sents,
            ))
        return out

    def parse_text_ai(
        self,
        text: str,
        ai_api_key: str,
        ai_base_url: str,
        ai_model: str,
        text_is_path: bool = False,
        cut_mode: str = "auto",
        custom_notes: str = "",
    ) -> dict:
        """AI 智能解析预览。"""
        if text_is_path:
            def _pcb(stage, done, total):
                if stage == "ocr":
                    self._progress("ocr", done / max(1, total) * 0.5,
                                   f"OCR 识别第 {done + 1}/{total} 页")
            text = read_text_source(text, progress_cb=_pcb)
        materials = ai_parse_materials(
            text, ai_api_key, ai_base_url, ai_model,
            progress_cb=lambda msg: self._progress("ai", 0.6, msg),
            cut_mode=cut_mode, custom_notes=custom_notes,
        )
        return {"materials": self._dump_materials(materials)}

    def process(
        self,
        text: str,
        audio_path: str,
        out_dir: str,
        text_is_path: bool = False,
        title: Optional[str] = None,
        model_size: Optional[str] = None,
        use_whisperx: bool = False,
        language: Optional[str] = None,
        min_match: Optional[float] = None,
        use_ai: bool = False,
        ai_api_key: str = "",
        ai_base_url: str = "",
        ai_model: str = "",
        cut_mode: str = "auto",
        custom_notes: str = "",
        materials: Optional[list] = None,
    ) -> dict:
        self._cancel.clear()
        title = title or "精听课件"
        model_size = model_size or "small"
        language = language or "en"
        min_match_f = float(min_match or 0.8)

        # 优先使用用户已校对的材料（跳过文稿解析）
        if materials:
            self._progress("parse", 0.03, "使用已校对的材料")
            mats = self._materials_from_dicts(materials)
        else:
            def _pcb(stage, done, total):
                if stage == "ocr":
                    self._progress("ocr", done / max(1, total) * 0.05,
                                   f"OCR 识别第 {done + 1}/{total} 页")
            text = read_text_source(text, progress_cb=_pcb) if text_is_path else text
            if use_ai:
                self._progress("ai", 0.06, "AI 解析文稿")
                mats = ai_parse_materials(
                    text, ai_api_key, ai_base_url, ai_model,
                    progress_cb=lambda msg: self._progress("ai", 0.07, msg),
                    cut_mode=cut_mode, custom_notes=custom_notes,
                )
            else:
                self._progress("parse", 0.06, "解析文本")
                mats = parse_paper_text(text)
        if not mats:
            raise RuntimeError("未在文本中识别到英文听力原文，请检查文本内容或改用 AI 智能解析")

        engine = self._ensure_asr(model_size)
        self._progress("asr", 0.1, "开始语音识别")
        tr = engine.transcribe(
            audio_path,
            language=language,
            use_whisperx=use_whisperx,
            cancel_check=self._cancel.is_set,
            progress_cb=lambda p, msg: self._progress("asr", 0.1 + p * 0.6, msg),
        )

        self._progress("align", 0.72, "对齐文本与音频")
        spans = align_materials(mats, tr.words, tr.duration, min_match=min_match_f)

        self._progress("cut", 0.8, "切分音频")
        json_path = export_courseware(
            spans, audio_path, out_dir, tr.duration,
            title=title, make_clips=True,
        )
        with open(json_path, "r", encoding="utf-8") as f:
            courseware = json.load(f)

        self._progress("done", 1.0, "完成")
        return {"courseware_path": json_path, "courseware": courseware}

    def re_export(self, out_dir: str, courseware: dict, audio_path: str = "") -> dict:
        """按（可能手动调整过的）课件时间重新切分导出切片与 courseware.json。"""
        src = audio_path or os.path.join(out_dir, courseware.get("source_audio") or "source.mp3")
        if not os.path.isfile(src):
            raise RuntimeError(f"找不到源音频: {src}")
        spans: List[MaterialSpan] = []
        for m in courseware.get("materials", []):
            sents = [
                SentenceSpan(
                    text=s.get("text", ""), speaker=s.get("speaker"),
                    start=float(s.get("start", 0.0)), end=float(s.get("end", 0.0)),
                    matched=bool(s.get("matched", True)),
                )
                for s in m.get("sentences", [])
            ]
            spans.append(MaterialSpan(
                id=int(m.get("id", len(spans) + 1)),
                title=m.get("title", ""), text=m.get("text", ""),
                start=float(m.get("start", 0.0)), end=float(m.get("end", 0.0)),
                match_rate=float(m.get("match_rate", 1.0)),
                needs_review=bool(m.get("needs_review", False)),
                sentences=sents,
            ))
        if not spans:
            raise RuntimeError("课件中没有材料可导出")
        self._progress("cut", 0.2, "重新切分音频")
        json_path = export_courseware(
            spans, src, out_dir, float(courseware.get("duration") or 0.0),
            title=courseware.get("title") or "精听课件", make_clips=True,
        )
        with open(json_path, "r", encoding="utf-8") as f:
            cw = json.load(f)
        self._progress("done", 1.0, "导出完成")
        return {"courseware_path": json_path, "courseware": cw}


def _send_stdout(obj: dict):
    sys.stdout.write(json.dumps(obj, ensure_ascii=False) + "\n")
    sys.stdout.flush()


def serve():
    sidecar = Sidecar(_send_stdout)
    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            req = json.loads(line)
        except json.JSONDecodeError:
            continue
        rid = req.get("id")
        method = req.get("method")
        params = req.get("params") or {}
        try:
            fn = getattr(sidecar, method, None)
            if fn is None or method.startswith("_"):
                raise RuntimeError(f"未知方法: {method}")
            result = fn(**params)
            _send_stdout({"id": rid, "result": result})
        except InterruptedError as e:
            _send_stdout({"id": rid, "error": str(e) or "已取消"})
        except Exception as e:
            traceback.print_exc(file=sys.stderr)
            _send_stdout({"id": rid, "error": f"{type(e).__name__}: {e}"})


def cli():
    ap = argparse.ArgumentParser(description="听力音频智能切分引擎 CLI")
    ap.add_argument("--text", required=True, help="试卷文本内容或 .txt/.docx 路径")
    ap.add_argument("--audio", required=True, help="原始听力音频路径")
    ap.add_argument("--out", required=True, help="课件输出目录")
    ap.add_argument("--title", default="精听课件")
    ap.add_argument("--model", default="small", help="whisper 模型：tiny/base/small/medium")
    ap.add_argument("--whisperx", action="store_true", help="启用 WhisperX 精对齐（需安装 whisperx）")
    ap.add_argument("--language", default="en")
    args = ap.parse_args()

    text = read_text_source(args.text)
    materials = parse_paper_text(text)
    print(f"[parse] 识别到 {len(materials)} 段英文材料", file=sys.stderr)
    for m in materials:
        print(f"  - {m.title}: {len(m.text)} 字符, {len(m.sentences)} 句", file=sys.stderr)
    if not materials:
        sys.exit("未识别到英文听力原文")

    print(f"[asr] 加载模型 {args.model} 并转录 {args.audio} ...", file=sys.stderr)
    engine = ASREngine(model_size=args.model, model_dir=os.environ.get("YP_MODEL_DIR"))
    tr = engine.transcribe(
        args.audio, language=args.language, use_whisperx=args.whisperx,
        progress_cb=lambda p, msg: print(f"[asr] {p*100:.0f}% {msg}", file=sys.stderr),
    )
    print(f"[asr] 转录完成：{len(tr.words)} 词, 时长 {tr.duration:.1f}s", file=sys.stderr)

    print("[align] 对齐中...", file=sys.stderr)
    spans = align_materials(materials, tr.words, tr.duration)
    for sp in spans:
        flag = " ⚠需人工确认" if sp.needs_review else ""
        print(f"  - {sp.title}: {sp.start:.1f}s ~ {sp.end:.1f}s (匹配率 {sp.match_rate:.0%}){flag}", file=sys.stderr)

    print("[cut] 切分并导出课件...", file=sys.stderr)
    json_path = export_courseware(spans, args.audio, args.out, tr.duration, title=args.title)
    print(f"[done] 课件已生成: {json_path}", file=sys.stderr)


if __name__ == "__main__":
    if "--cli" in sys.argv:
        sys.argv.remove("--cli")
        cli()
    else:
        serve()
